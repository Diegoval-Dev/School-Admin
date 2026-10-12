from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from openpyxl import load_workbook

from colegio.models import Asistencia, CicloEscolar, EstadoAsistencia, Estudiante, Genero, Grado, Jornada, Matricula
from colegio.services.asistencia import ciclar_asistencia, dias_del_mes, filas_de_asistencia


class AsistenciaTestCase(TestCase):
    def setUp(self):
        self.ciclo = CicloEscolar.objects.create(anio=2026, activo=True)
        self.matricula = self._crear_matricula("Pérez", "Ana")
        self.usuario = get_user_model().objects.create_superuser("directora", "", "clave-muy-segura-2026")
        self.client.force_login(self.usuario)

    def _crear_matricula(self, apellido, nombre, jornada=Jornada.DIARIO, activa=True):
        estudiante = Estudiante.objects.create(primer_apellido=apellido, primer_nombre=nombre, genero=Genero.FEMENINO)
        return Matricula.objects.create(
            estudiante=estudiante, ciclo=self.ciclo, grado=Grado.PRIMERO, jornada=jornada, activa=activa
        )


class DiasDelMesTests(TestCase):
    def test_diario_incluye_todos_los_dias_del_mes(self):
        self.assertEqual(len(dias_del_mes(2026, 2, Jornada.DIARIO)), 28)

    def test_fin_de_semana_solo_sabados_y_domingos(self):
        dias = dias_del_mes(2026, 1, Jornada.FIN_SEMANA)

        self.assertTrue(all(dia.weekday() >= 5 for dia in dias))
        self.assertEqual(len(dias), 9)


class CiclarAsistenciaTests(AsistenciaTestCase):
    def test_primer_clic_marca_presente(self):
        ciclar_asistencia(self.matricula, date(2026, 1, 5))

        self.assertEqual(Asistencia.objects.get().estado, EstadoAsistencia.PRESENTE)

    def test_ciclo_completo_de_estados_sin_borrar_filas(self):
        dia = date(2026, 1, 5)
        estados_por_clic = []
        for _ in range(6):
            estados_por_clic.append(ciclar_asistencia(self.matricula, dia).estado)

        self.assertEqual(
            estados_por_clic,
            [
                EstadoAsistencia.PRESENTE,
                EstadoAsistencia.AUSENTE,
                EstadoAsistencia.EXCUSA,
                EstadoAsistencia.SIN_REGISTRO,
                EstadoAsistencia.PRESENTE,
                EstadoAsistencia.AUSENTE,
            ],
        )
        self.assertEqual(Asistencia.objects.filter(matricula=self.matricula, fecha=dia).count(), 1)

    def test_sin_registro_se_guarda_como_fila(self):
        dia = date(2026, 1, 5)
        for _ in range(4):
            ciclar_asistencia(self.matricula, dia)

        self.assertTrue(Asistencia.objects.filter(estado=EstadoAsistencia.SIN_REGISTRO).exists())

    def test_cada_cambio_queda_en_el_historial(self):
        dia = date(2026, 1, 5)
        ciclar_asistencia(self.matricula, dia)
        ciclar_asistencia(self.matricula, dia)

        asistencia = Asistencia.objects.get()
        self.assertEqual(asistencia.history.count(), 2)
        self.assertEqual(asistencia.history.order_by("history_id").first().estado, EstadoAsistencia.PRESENTE)

    def test_una_asistencia_por_matricula_y_dia(self):
        dia = date(2026, 1, 5)
        ciclar_asistencia(self.matricula, dia)
        ciclar_asistencia(self.matricula, dia)

        self.assertEqual(Asistencia.objects.filter(matricula=self.matricula, fecha=dia).count(), 1)

    def test_matricula_inactiva_no_recibe_asistencia(self):
        inactiva = self._crear_matricula("Lara", "Luis", activa=False)

        with self.assertRaises(ValidationError):
            ciclar_asistencia(inactiva, date(2026, 1, 5))

    def test_filas_solo_incluyen_matriculas_activas_del_filtro(self):
        self._crear_matricula("Lara", "Luis", activa=False)
        self._crear_matricula("Gómez", "Eva", jornada=Jornada.FIN_SEMANA)
        activas = Matricula.objects.filter(activa=True, jornada=Jornada.DIARIO)

        filas = filas_de_asistencia(activas, dias_del_mes(2026, 1, Jornada.DIARIO))

        self.assertEqual([fila["matricula"] for fila in filas], [self.matricula])


class SinRegistroTests(AsistenciaTestCase):
    def test_sin_registro_no_cuenta_en_totales(self):
        ciclar_asistencia(self.matricula, date(2026, 1, 5))
        for _ in range(2):
            ciclar_asistencia(self.matricula, date(2026, 1, 6))
        for _ in range(4):
            ciclar_asistencia(self.matricula, date(2026, 1, 7))

        fila = filas_de_asistencia(
            Matricula.objects.filter(pk=self.matricula.pk), dias_del_mes(2026, 1, Jornada.DIARIO)
        )[0]

        self.assertEqual((fila["presentes"], fila["ausentes"], fila["excusas"]), (1, 1, 0))

    def test_sin_registro_se_ve_en_blanco_en_reporte_y_excel(self):
        for _ in range(4):
            ciclar_asistencia(self.matricula, date(2026, 1, 5))

        reporte = self.client.get("/asistencia/reporte/", {"ciclo": self.ciclo.pk, "mes": 1})
        excel = self.client.get("/asistencia/excel/", {"ciclo": self.ciclo.pk, "mes": 1})

        hoja = load_workbook(BytesIO(excel.content)).active
        dia_cinco = hoja.cell(row=2, column=3 + 5).value
        self.assertEqual(reporte.status_code, 200)
        self.assertIsNone(dia_cinco)


class TotalesAsistenciaTests(AsistenciaTestCase):
    def test_totales_por_estudiante(self):
        ciclar_asistencia(self.matricula, date(2026, 1, 5))
        ciclar_asistencia(self.matricula, date(2026, 1, 6))
        ciclar_asistencia(self.matricula, date(2026, 1, 6))
        ciclar_asistencia(self.matricula, date(2026, 1, 7))
        ciclar_asistencia(self.matricula, date(2026, 1, 7))
        ciclar_asistencia(self.matricula, date(2026, 1, 7))

        fila = filas_de_asistencia(
            Matricula.objects.filter(pk=self.matricula.pk), dias_del_mes(2026, 1, Jornada.DIARIO)
        )[0]

        self.assertEqual((fila["presentes"], fila["ausentes"], fila["excusas"]), (1, 1, 1))


class VistasAsistenciaTests(AsistenciaTestCase):
    def test_tabla_responde_con_el_estudiante(self):
        respuesta = self.client.get("/asistencia/", {"ciclo": self.ciclo.pk, "mes": 1})

        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "Pérez")

    def test_clic_en_celda_cambia_estado_y_redirige(self):
        respuesta = self.client.post(
            "/asistencia/ciclar/", {"celda": f"{self.matricula.pk}|2026-01-05"}
        )

        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(Asistencia.objects.get().estado, EstadoAsistencia.PRESENTE)

    def test_celda_invalida_responde_404(self):
        respuesta = self.client.post("/asistencia/ciclar/", {"celda": "basura"})

        self.assertEqual(respuesta.status_code, 404)

    def test_excel_y_reporte_responden(self):
        excel = self.client.get("/asistencia/excel/", {"ciclo": self.ciclo.pk, "mes": 1})
        reporte = self.client.get("/asistencia/reporte/", {"ciclo": self.ciclo.pk, "mes": 1})

        self.assertTrue(excel.content.startswith(b"PK"))
        self.assertEqual(reporte.status_code, 200)
